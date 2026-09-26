"""Strict reader for the two cold-start experiments in attachment 2."""

from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXPERIMENT_FILE = ROOT / "附件" / "附件2.xlsx"


@dataclass(frozen=True)
class Experiment:
    label: str
    ambient_c: float
    time_s: np.ndarray
    current_a: np.ndarray
    current_density_acm2: np.ndarray
    voltage_v: np.ndarray
    temperature_c: np.ndarray


def load_experiments(path: Path = EXPERIMENT_FILE) -> dict[str, Experiment]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    result = {}
    for label, ambient in [('-20℃', -20.0), ('-25℃', -25.0)]:
        ws = workbook[label]
        rows = []
        for row in ws.iter_rows(min_row=3, max_col=5, values_only=True):
            if all(value is None for value in row):
                continue
            if not all(isinstance(value, (int, float)) for value in row):
                raise ValueError(f'{label}: nonnumeric measurement row: {row}')
            rows.append(tuple(float(value) for value in row))
        values = np.asarray(rows, dtype=float)
        if values.shape != (184, 5):
            raise ValueError(f'{label}: expected 184x5 measurements, got {values.shape}')
        if not np.all(np.isfinite(values)) or not np.all(np.diff(values[:, 0]) > 0):
            raise ValueError(f'{label}: missing, nonfinite, or duplicate time values')
        if abs(values[0, 0]) > 1e-9 or abs(values[-1, 0] - 36.6) > 1e-8:
            raise ValueError(f'{label}: unexpected experiment time range')
        result[label] = Experiment(
            label=label,
            ambient_c=ambient,
            time_s=values[:, 0],
            current_a=values[:, 1],
            voltage_v=values[:, 2],
            temperature_c=values[:, 3],
            current_density_acm2=values[:, 4],
        )
    return result
